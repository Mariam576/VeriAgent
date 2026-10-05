"""Validation, text extraction and chunking - no database needed."""

from __future__ import annotations

import io

import docx
import pymupdf
import pytest

from app.errors import DocumentProcessingError, FileTooLarge, UnsupportedFileType, ValidationFailed
from app.rag.chunking import chunk_pages
from app.rag.loaders import PageText, clean_text, guess_document_type, guess_year, load_pages
from app.security.validation import (
    sanitize_filename,
    validate_document_type,
    validate_question,
    validate_upload,
)


def make_pdf(pages: list[str]) -> bytes:
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


# -------------------------------------------------------------- validation
def test_sanitize_filename_removes_directories_and_odd_characters():
    assert sanitize_filename("../../etc/passwd") == "passwd"
    assert sanitize_filename("C:\\Users\\me\\Invoice 102.pdf") == "Invoice 102.pdf"
    assert sanitize_filename("in<voice>;rm -rf.pdf") == "in_voice__rm -rf.pdf"
    assert sanitize_filename("") == "document"
    assert sanitize_filename(None) == "document"


def test_allowed_file_passes(settings):
    name, extension = validate_upload("Contract.pdf", make_pdf(["hello"]), settings)
    assert (name, extension) == ("Contract.pdf", ".pdf")


@pytest.mark.parametrize("filename", ["malware.exe", "script.js", "archive.zip", "noextension"])
def test_disallowed_extensions_are_rejected(settings, filename):
    with pytest.raises(UnsupportedFileType):
        validate_upload(filename, b"data", settings)


def test_extension_must_match_real_content(settings):
    with pytest.raises(UnsupportedFileType, match="not a PDF"):
        validate_upload("fake.pdf", b"MZ this is really an executable", settings)
    with pytest.raises(UnsupportedFileType, match="not a Word"):
        validate_upload("fake.docx", b"plain text", settings)
    with pytest.raises(UnsupportedFileType, match="binary"):
        validate_upload("fake.txt", b"abc\x00\x01\x02", settings)


def test_empty_and_oversized_files_are_rejected(settings):
    with pytest.raises(ValidationFailed, match="empty"):
        validate_upload("a.txt", b"", settings)
    too_big = b"a" * (settings.max_upload_bytes + 1)
    with pytest.raises(FileTooLarge):
        validate_upload("a.txt", too_big, settings)


def test_document_type_must_be_known():
    assert validate_document_type(None) is None
    assert validate_document_type(" Purchase Order ") == "purchase_order"
    with pytest.raises(ValidationFailed):
        validate_document_type("love_letter")


def test_question_validation(settings):
    assert validate_question("  Compare   the\tinvoice\x00 ", settings) == "Compare the invoice"
    with pytest.raises(ValidationFailed, match="empty"):
        validate_question("   ", settings)
    with pytest.raises(ValidationFailed, match="longer"):
        validate_question("x" * (settings.MAX_QUESTION_CHARS + 1), settings)


# ----------------------------------------------------------------- loaders
def test_pdf_pages_keep_their_numbers():
    pages = load_pages(make_pdf(["First page text", "Second page text", "Third page text"]), ".pdf")
    assert [page.page_number for page in pages] == [1, 2, 3]
    assert "Second page" in pages[1].text


def test_sample_contract_has_value_and_terms_on_page_4(samples_dir):
    pages = load_pages((samples_dir / "Contract.pdf").read_bytes(), ".pdf")
    assert len(pages) == 5
    assert "£40,000" in pages[3].text
    assert "within 30 days" in pages[3].text
    assert "£40,000" not in pages[0].text


def test_docx_text_and_tables_are_extracted():
    document = docx.Document()
    document.add_paragraph("Purchase order PO-4421")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Total"
    table.rows[0].cells[1].text = "£58,000"
    buffer = io.BytesIO()
    document.save(buffer)

    pages = load_pages(buffer.getvalue(), ".docx")
    assert len(pages) == 1 and pages[0].page_number == 1
    assert "PO-4421" in pages[0].text
    assert "Total | £58,000" in pages[0].text


def test_txt_form_feed_starts_a_new_page():
    pages = load_pages("page one\fpage two".encode(), ".txt")
    assert [(page.page_number, page.text) for page in pages] == [(1, "page one"), (2, "page two")]


def test_bad_or_empty_documents_raise_a_clear_error():
    with pytest.raises(DocumentProcessingError, match="could not be opened|could not be read"):
        load_pages(b"%PDF-1.7 this is not really a pdf", ".pdf")
    with pytest.raises(DocumentProcessingError, match="No text could be extracted"):
        load_pages(make_pdf(["", ""]), ".pdf")  # e.g. a scanned document
    with pytest.raises(DocumentProcessingError, match="No loader"):
        load_pages(b"data", ".xyz")


def test_clean_text_normalises_whitespace():
    assert clean_text("A  lot\tof   space \n\n\n\n next\x00") == "A lot of space\n\nnext"


def test_guess_document_type_and_year():
    assert guess_document_type("Invoice_102.pdf", "") == "invoice"
    assert guess_document_type("Procurement_Policy.pdf", "") == "policy"
    assert guess_document_type("scan0001.pdf", "SERVICES AGREEMENT between ...") == "contract"
    assert guess_document_type("scan0001.pdf", "nothing recognisable") == "other"
    assert guess_year("Signed 12 January 2026. Renewal due 2027. Dated 2026.") == 2026
    assert guess_year("no year here") is None


# ---------------------------------------------------------------- chunking
def test_chunks_respect_size_and_keep_page_numbers():
    long_text = " ".join(f"Sentence number {i} about maintenance." for i in range(120))
    pages = [PageText(1, "Short first page."), PageText(2, long_text), PageText(3, "")]
    chunks = chunk_pages(pages, chunk_size=300, chunk_overlap=50)

    assert all(len(chunk.text) <= 300 for chunk in chunks)
    assert chunks[0].page_number == 1
    assert {chunk.page_number for chunk in chunks[1:]} == {2}  # empty page 3 gives no chunk
    assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))
    assert len(chunks) > 5


def test_neighbouring_chunks_overlap():
    text = " ".join(f"word{i}" for i in range(400))
    chunks = chunk_pages([PageText(1, text)], chunk_size=200, chunk_overlap=60)
    first_words = set(chunks[0].text.split())
    second_words = set(chunks[1].text.split())
    assert first_words & second_words  # shared text at the boundary
